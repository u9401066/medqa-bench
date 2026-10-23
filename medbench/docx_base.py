"""Small python-docx layer for the report: CJK-aware styles, inline Markdown (bold, code, links), tables, code blocks,
figures with captions, a static table of contents and page-numbered footers."""
from __future__ import annotations

import io
import re
from typing import List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LATIN = "Calibri"
CJK = "Microsoft JhengHei"
MONO = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "DCE6F1"
CODE_FILL = "F2F2F2"
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\))")
ALIGN = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}


def _fonts(rpr, latin: str = LATIN, east: str = CJK) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), latin)
    rfonts.set(qn("w:eastAsia"), east)


def _shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _field(paragraph, instr: str, size: float) -> None:
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("text", "1"), ("end", None)):
        r = paragraph.add_run()
        r.font.size = Pt(size)
        if kind in ("begin", "separate", "end"):
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif kind == "text":
            r.text = text
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = " %s " % text
            r._r.append(it)


class DocBuilder:
    def __init__(self, header_text: str):
        self.d = Document()
        self.fig_n = 0
        self.tab_n = 0
        self.headings: List[Tuple[int, str]] = []
        self._toc_marker = None
        self._setup(header_text)

    # ---- page and styles
    def _setup(self, header_text: str) -> None:
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
        st = self.d.styles
        normal = st["Normal"]
        normal.font.size = Pt(10)
        _fonts(normal.element.get_or_add_rPr())
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.15
        for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11), ("Caption", 9)):
            s = st[name]
            _fonts(s.element.get_or_add_rPr())
            s.font.size = Pt(size)
            s.font.color.rgb = ACCENT if name != "Caption" else GREY
            if name.startswith("Heading"):
                s.font.bold = True
                s.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
                s.paragraph_format.space_after = Pt(6)
                s.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Bullet 2", "List Bullet 3"):
            _fonts(st[name].element.get_or_add_rPr())
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(header_text)
        r.font.size, r.font.color.rgb = Pt(8), GREY
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("第 ")
        r.font.size = Pt(8)
        _field(fp, "PAGE", 8)
        r = fp.add_run(" 頁，共 ")
        r.font.size = Pt(8)
        _field(fp, "NUMPAGES", 8)
        r = fp.add_run(" 頁")
        r.font.size = Pt(8)

    # ---- inline text
    def link(self, p, url: str, text: str, size: Optional[float] = None) -> None:
        r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "0563C1")
        rpr.append(c)
        if size:
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), str(int(round(size * 2))))
            rpr.append(sz)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        h.append(r)
        p._p.append(h)

    def inline(self, p, text: str, size: Optional[float] = None, bold: bool = False, color: Optional[RGBColor] = None) -> None:
        lines = text.split("\n")
        for li, line in enumerate(lines):
            for part in INLINE.split(line):
                if not part:
                    continue
                if part.startswith("**") and part.endswith("**") and len(part) > 4:
                    self.inline(p, part[2:-2], size, True, color)
                    continue
                m = re.fullmatch(r"\[([^\]]+)\]\(([^)\s]+)\)", part)
                if m:
                    self.link(p, m.group(2), m.group(1), size)
                    continue
                code = part.startswith("`") and part.endswith("`") and len(part) > 2
                r = p.add_run(part[1:-1] if code else part)
                if code:
                    r.font.name = MONO
                    _fonts(r._r.get_or_add_rPr(), MONO, CJK)
                    r.font.size = Pt((size or 10) - 0.5)
                elif size:
                    r.font.size = Pt(size)
                if bold:
                    r.bold = True
                if color is not None:
                    r.font.color.rgb = color
            if li < len(lines) - 1:
                p.add_run().add_break()

    # ---- blocks
    def heading(self, text: str, level: int = 1) -> None:
        self.d.add_heading(text, level)
        self.headings.append((level, text))

    def para(self, text: str = "", size: Optional[float] = None, align: str = "l", bold: bool = False,
             color: Optional[RGBColor] = None, style: Optional[str] = None, space_after: Optional[float] = None):
        p = self.d.add_paragraph(style=style)
        p.alignment = ALIGN[align]
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        self.inline(p, text, size, bold, color)
        return p

    def note(self, text: str) -> None:
        self.para(text, size=8.5, color=GREY)

    def bullet(self, text: str, level: int = 0) -> None:
        p = self.d.add_paragraph(style=("List Bullet", "List Bullet 2", "List Bullet 3")[min(level, 2)])
        p.paragraph_format.space_after = Pt(2)
        self.inline(p, text)

    def numbered(self, number: str, text: str, level: int = 0) -> None:
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.9 + 0.7 * level)
        pf.first_line_indent = Cm(-0.6)
        pf.space_after = Pt(2)
        self.inline(p, "%s %s" % (number, text))

    def caption(self, text: str) -> None:
        p = self.d.add_paragraph(style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        self.inline(p, text)

    def table(self, header: Sequence[str], rows: Sequence[Sequence[object]], align: Optional[str] = None,
              widths: Optional[Sequence[float]] = None, size: float = 8.5, caption: Optional[str] = None,
              note: Optional[str] = None):
        if caption:
            self.tab_n += 1
            self.caption("表 %d：%s" % (self.tab_n, caption))
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        align = align or "l" * len(header)
        for i, h in enumerate(header):
            self._cell(t.rows[0].cells[i], str(h), size, "c", bold=True)
            _shade(t.rows[0].cells[i], HEADER_FILL)
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                self._cell(cells[i], "" if v is None else str(v), size, align[i])
        if widths:
            t.autofit = False
            for i, w in enumerate(widths):
                t.columns[i].width = Cm(w)
                for row in t.rows:
                    row.cells[i].width = Cm(w)
        if note:
            self.note(note)
        else:
            self.para("", size=4, space_after=2)
        return t

    def merge_row(self, t, row: int, first: int, last: int, text: str, size: float = 8.5, align: str = "c") -> None:
        cell = t.rows[row].cells[first].merge(t.rows[row].cells[last])
        for extra in cell.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        for r in list(cell.paragraphs[0].runs):
            r._r.getparent().remove(r._r)
        self._cell(cell, text, size, align)

    def _cell(self, cell, text: str, size: float, align: str, bold: bool = False) -> None:
        p = cell.paragraphs[0]
        p.alignment = ALIGN[align]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        self.inline(p, text, size, bold)

    def code(self, text: str, size: float = 7.5) -> None:
        t = self.d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        cell = t.rows[0].cells[0]
        _shade(cell, CODE_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        lines = text.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            r = p.add_run(line)
            r.font.name = MONO
            _fonts(r._r.get_or_add_rPr(), MONO, CJK)
            r.font.size = Pt(size)
            if i < len(lines) - 1:
                r.add_break()
        self.para("", size=4, space_after=2)

    def figure(self, png: bytes, caption: str, width_cm: float = 16.5) -> None:
        self.fig_n += 1
        self.d.add_picture(io.BytesIO(png), width=Cm(width_cm))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.d.paragraphs[-1].paragraph_format.keep_with_next = True
        self.caption("圖 %d：%s" % (self.fig_n, caption))

    def page_break(self) -> None:
        self.d.add_page_break()

    # ---- table of contents (static: works in every viewer without updating fields)
    def toc_here(self) -> None:
        self._toc_marker = self.d.add_paragraph()

    def finish_toc(self, max_level: int = 2) -> None:
        if self._toc_marker is None:
            return
        for level, text in self.headings:
            if level > max_level:
                continue
            p = self._toc_marker.insert_paragraph_before()
            p.paragraph_format.left_indent = Cm(0.8 * (level - 1))
            p.paragraph_format.space_after = Pt(1 if level > 1 else 3)
            r = p.add_run(text)
            r.font.size = Pt(10.5 if level == 1 else 9.5)
            r.bold = level == 1
        self._toc_marker._p.getparent().remove(self._toc_marker._p)
        self._toc_marker = None

    # ---- simple Markdown (the report's notes file): ##/### headings, bullets, numbered items, paragraphs
    def markdown(self, text: str) -> None:
        for line in text.splitlines():
            if not line.strip() or line.strip().startswith("<!--"):
                continue
            if line.startswith("### "):
                self.heading(line[4:].strip(), 2)
            elif line.startswith("## "):
                self.heading(line[3:].strip(), 1)
            elif re.match(r"^\s*- ", line):
                self.bullet(line.strip()[2:], (len(line) - len(line.lstrip())) // 2)
            elif re.match(r"^\s*\d+\. ", line):
                m = re.match(r"^(\s*)(\d+\.) (.*)$", line)
                self.numbered(m.group(2), m.group(3), len(m.group(1)) // 2)
            else:
                self.para(line.strip())

    def save(self, path) -> None:
        self.finish_toc()
        self.d.save(str(path))
